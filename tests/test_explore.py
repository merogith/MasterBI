"""Portfolio workflow: calculations, persistence, exports and AI boundaries."""

import io
import json
from types import SimpleNamespace

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pypdf import PdfReader

from kpi_maker.ai.client import AIUnavailable, Call, Usage, use_factory
from kpi_maker.api.server import app
from kpi_maker.explore import api as workspace
from kpi_maker.explore.engine import Analysis, Filter, Measure, compute
from kpi_maker.explore.exports import export
from kpi_maker.ingest.pipeline import derive_pl_columns
from kpi_maker.survey.engine import build_profile, random_answers


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(workspace, "root", lambda: tmp_path)
    return TestClient(app)


def test_unknown_data_keeps_dates_and_ids(client):
    content = b"date,player_id,score\n2025-01-01,001,4\n2025-01-02,002,6\n"
    response = client.post(
        "/api/explore/upload", files=[("files", ("games.csv", content, "text/csv"))]
    )
    assert response.status_code == 200, response.text
    p = response.json()
    assert p["spec"]["measures"][0]["aggregation"] == "count"
    raw = workspace.tables_for(workspace.load(p["id"]))["table_1"]
    assert raw["date"].tolist() == ["2025-01-01", "2025-01-02"]
    assert raw["player_id"].tolist() == ["001", "002"]
    assert (
        client.get(f"/api/explore/projects/{p['id']}").json()["result"] == p["result"]
    )


def test_all_workbook_sheets_and_file_limit(client, monkeypatch):
    stream = io.BytesIO()
    with pd.ExcelWriter(stream, engine="openpyxl") as book:
        pd.DataFrame({"period": [1, 2]}).to_excel(
            book, sheet_name="Health", index=False
        )
        pd.DataFrame({"wins": [3, 4]}).to_excel(book, sheet_name="Games", index=False)
    r = client.post(
        "/api/explore/upload", files=[("files", ("data.xlsx", stream.getvalue()))]
    )
    assert r.status_code == 200, r.text
    assert len(r.json()["tables"]) == 2
    monkeypatch.setattr(workspace, "MAX_FILE_BYTES", 10)
    r = client.post(
        "/api/explore/upload", files=[("files", ("large.csv", b"field\n1234567890\n"))]
    )
    assert r.status_code == 422


def test_weighted_ratio_and_missing_are_not_zero():
    tables = {
        "t": pd.DataFrame(
            {
                "team": ["A", "A", "B", "B"],
                "wins": ["1", "9", "0", "bad"],
                "games": ["1", "10", "0", "2"],
            }
        )
    }
    spec = Analysis(
        table="t",
        dimension="team",
        measures=[
            Measure(
                name="Rate", column="wins", aggregation="ratio", denominator="games"
            )
        ],
    )
    r = compute(spec, tables)
    assert r["totals"][0]["value"] == pytest.approx(10 / 11)
    assert r["rows"][1]["values"] == [None]
    assert any("nonnumeric" in w for w in r["warnings"])
    spec.measures = [Measure(name="sum", column="wins", aggregation="sum")]
    spec.filters = [Filter(column="team", value="C")]
    assert compute(spec, tables)["totals"][0]["value"] is None


def test_cleaning_filters_and_undo(client):
    r = client.post(
        "/api/explore/upload",
        files=[("files", ("test.csv", b"category,value\n A ,10\nA,10\nB,-2\n"))],
    )
    p = r.json()
    url = f"/api/explore/projects/{p['id']}"
    spec = p["spec"] | {
        "trim_text": True,
        "drop_duplicates": True,
        "dimension": "category",
        "measures": [
            {
                "name": "Total",
                "column": "value",
                "aggregation": "sum",
                "denominator": "",
            }
        ],
    }
    r = client.put(url, json={"spec": spec, "revision": 0})
    assert r.status_code == 200, r.text
    assert r.json()["result"]["totals"][0]["value"] == 8
    assert r.json()["result"]["included_rows"] == 2
    assert client.put(url, json={"spec": spec, "revision": 0}).status_code == 409
    undone = client.post(url + "/undo").json()
    assert undone["spec"] == p["spec"]
    assert undone["result"]["included_rows"] == 3
    bad = spec | {"dimension": "invented"}
    assert client.put(url, json={"spec": bad, "revision": 2}).status_code == 422


@pytest.mark.parametrize("kind", ["retail", "saas", "pokemon"])
def test_demos_export_same_numbers(client, kind):
    r = client.post("/api/explore/demo/" + kind)
    assert r.status_code == 200, r.text
    p = r.json()
    url = f"/api/explore/projects/{p['id']}"
    pdf = client.get(url + "/export/pdf")
    assert pdf.status_code == 200, pdf.text[:100] if pdf.status_code != 200 else ""
    text = "\n".join(
        page.extract_text() for page in PdfReader(io.BytesIO(pdf.content)).pages
    )
    assert "Synthetic demonstration" in text
    assert p["spec"]["measures"][0]["name"] in text
    deck = Presentation(io.BytesIO(client.get(url + "/export/pptx").content))
    charts = [
        shape.chart
        for slide in deck.slides
        for shape in slide.shapes
        if shape.has_chart
    ]
    assert charts
    assert list(charts[0].series[0].values) == pytest.approx(
        [row["values"][0] for row in p["result"]["rows"]]
    )
    assert client.get(url + "/export/html").status_code == 200
    assert client.get(url + "/export/csv").status_code == 200
    if kind == "pokemon":
        assert p["result"]["totals"][0]["value"] == 0.5
        assert p["result"]["totals"][1]["value"] == 160


def test_escaping_and_spreadsheet_formula():
    spec = Analysis(table="t", title="<script>bad()</script>", dimension="category")
    result = compute(
        spec, {"t": pd.DataFrame({"category": ['=HYPERLINK("x")', "<script>"]})}
    )
    html, _ = export(spec, result, "html")
    assert b"<script>" not in html
    csv, _ = export(spec, result, "csv")
    assert "'=HYPERLINK" in csv.decode("utf-8-sig")


def test_chat_persists_and_cannot_apply_invalid_plan(client):
    p = client.post("/api/explore/demo/pokemon").json()
    url = f"/api/explore/projects/{p['id']}"

    class Fake:
        calls = []

        def json(self, **kwargs):
            self.calls = [
                Call("chat", "gpt-5.6-luna", Usage(input_tokens=100, output_tokens=200))
            ]
            proposed = p["spec"] | {"dimension": "nonexistent"}
            return {
                "message": "Try this grouping.",
                "analysis_json": json.dumps(proposed),
            }

    old = use_factory(lambda _: Fake())
    try:
        r = client.post(url + "/chat", json={"message": "Change the grouping"})
    finally:
        use_factory(old)
    assert r.status_code == 200, r.text
    saved = client.get(url).json()
    assert len(saved["messages"]) == 3
    assert saved["proposal"] is None
    assert saved["spec"] == p["spec"]
    assert saved["spent_usd"] > 0


def test_chat_budget_blocks_before_client(client):
    p = client.post("/api/explore/demo/pokemon").json()
    old = use_factory(lambda _: pytest.fail("Paid client must not be constructed"))
    try:
        r = client.post(
            f"/api/explore/projects/{p['id']}/chat",
            json={"message": "Explain", "tier": "advanced", "budget_usd": 0.01},
        )
    finally:
        use_factory(old)
    assert r.status_code == 422


def test_chat_proposal_recalculates_exports_and_can_undo(client):
    p = client.post("/api/explore/demo/retail").json()
    url = f"/api/explore/projects/{p['id']}"

    class Fake:
        calls = []

        def json(self, **kwargs):
            return {
                "message": "Show average sales.",
                "analysis_json": json.dumps(
                    p["spec"]
                    | {
                        "measures": [
                            {
                                "name": "Average sale",
                                "column": "sales",
                                "aggregation": "mean",
                                "denominator": "",
                            }
                        ]
                    }
                ),
            }

    old = use_factory(lambda _: Fake())
    try:
        proposed = client.post(
            url + "/chat", json={"message": "Use average sales"}
        ).json()
    finally:
        use_factory(old)
    assert proposed["proposal"]
    changed = client.put(url, json={"spec": proposed["proposal"], "revision": 0}).json()
    assert changed["result"]["totals"][0]["value"] == pytest.approx(p["result"]["totals"][0]["value"] / 360)
    assert client.post(url + "/undo").json()["spec"] == p["spec"]


def test_origin_and_file_boundary(client, tmp_path, monkeypatch):
    from kpi_maker.api import server

    monkeypatch.setattr(server, "RUNS_DIR", tmp_path)
    (tmp_path / "abc").mkdir()
    (tmp_path / "abcdef").mkdir()
    (tmp_path / "abcdef" / "report.pdf").write_bytes(b"private")
    with pytest.raises(Exception) as exc:
        server.serve_file("abc", "../abcdef/report.pdf")
    assert exc.value.status_code == 404
    assert (
        client.post(
            "/api/explore/demo/retail", headers={"origin": "https://untrusted.example"}
        ).status_code
        == 403
    )
    assert (
        client.get(
            "/api/health", headers={"origin": "http://localhost:5173"}
        ).status_code
        == 200
    )


def test_partial_expenses_and_zero_margin():
    frame, _ = derive_pl_columns(
        pd.DataFrame({"revenue": [1000.0], "cogs": [400.0], "marketing_cost": [100.0]})
    )
    assert "ebitda" not in frame and "total_opex" not in frame
    frame, _ = derive_pl_columns(pd.DataFrame({"revenue": [0.0], "cogs": [0.0]}))
    assert pd.isna(frame["gross_margin_pct"].iloc[0])


def test_survey_diversity_and_open_multi_answers():
    assert len({random_answers(s)["business_model"] for s in range(50)}) > 3
    profile = build_profile(
        {
            "secondary": ["margin_expansion", "margin_expansion"],
            "question": "Where are the risks?",
        }
    )
    assert len(profile.intent.secondary) == 1
    assert profile.intent.question == "Where are the risks?"
